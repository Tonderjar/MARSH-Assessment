import os
import glob
from datetime import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference


REQUIRED_COLS = {"Client", "Country", "Currency", "Transaction"}

############################################################################################
#                           File reading and parsing logic                                 #
############################################################################################

def _find_header_row(raw: pd.DataFrame) -> int | None:
    """Return the integer row index whose values contain all four required column names."""
    for i, row in raw.iterrows():
        if REQUIRED_COLS.issubset({str(v).strip() for v in row}):
            return i
    return None


def _read_file(filepath: str) -> pd.DataFrame | None:
    ext = os.path.splitext(filepath)[1].lower()
    try:
        if ext == ".csv":
            raw = pd.read_csv(filepath, header=None, sep=None, engine="python", dtype=str)
        else:
            raw = pd.read_excel(filepath, header=None, dtype=str)
    except Exception as e:
        print(f"Skipping {filepath}: {e}")
        return None

    header_row = _find_header_row(raw)
    if header_row is None:
        print(f"Skipping {filepath}: required columns not found")
        return None

    raw.columns = [str(v).strip() for v in raw.iloc[header_row]]
    df = raw.iloc[header_row + 1:].reset_index(drop=True)
    df = df[list(REQUIRED_COLS)].copy()
    return df


def load_all_transactions(data_root: str) -> pd.DataFrame:
    """
    Walk the data folder tree and concatenate every DataTrans_*.csv / *.xlsx file
    into a single DataFrame.

    Expected layout:
        <data_root>/<year>/<NN Month Name>/<DD-MM-YY>/DataTrans_N.{csv,xlsx}

    Added column:
        log_date (datetime) — parsed from the DD-MM-YY folder name.
    """
    frames = []

    for dirpath, _, filenames in os.walk(data_root):
        for filename in sorted(filenames):
            if not filename.startswith("DataTrans_"):
                continue
            if os.path.splitext(filename)[1].lower() not in {".csv", ".xlsx"}:
                continue

            date_folder = os.path.basename(dirpath)
            try:
                log_date = datetime.strptime(date_folder, "%d-%m-%y")
            except ValueError:
                log_date = pd.NaT

            df = _read_file(os.path.join(dirpath, filename))
            if df is None:
                continue

            df["log_date"] = log_date
            frames.append(df)

    if not frames:
        return pd.DataFrame(columns=["Client", "Country", "Currency", "Transaction", "log_date"])

    return pd.concat(frames, ignore_index=True)


def load_fx_rates(fx_path: str) -> dict[str, float]:
    """
    Parse the FX Rates xlsx and return a dict mapping currency CODE to its RATE vs USD.
    Searches for the row containing CODE and RATE to skip any title rows at the top.
    """
    raw = pd.read_excel(fx_path, header=None, dtype=str)
    fx_cols = {"CODE", "RATE"}
    header_row = None
    for i, row in raw.iterrows():
        if fx_cols.issubset({str(v).strip().upper() for v in row}):
            header_row = i
            break
    if header_row is None:
        raise ValueError(f"Could not find CODE/RATE header in {fx_path}")

    raw.columns = [str(v).strip().upper() for v in raw.iloc[header_row]]
    df = raw.iloc[header_row + 1:].reset_index(drop=True)
    df = df[["CODE", "RATE"]].dropna(subset=["CODE", "RATE"])
    df["RATE"] = pd.to_numeric(df["RATE"], errors="coerce")
    rates = dict(zip(df["CODE"].str.strip(), df["RATE"]))
    rates.setdefault("USD", 1.0)
    return rates


############################################################################################
#                           Simple aggregation helpers (public)                            #
############################################################################################

def currency_client_transactions(df: pd.DataFrame) -> pd.DataFrame:
    """Count the transactions by Client and Currency."""
    return df.groupby(["Client", "Currency"]).size().reset_index(name="Count")


def client_totals_usd(df: pd.DataFrame, fx_rates: dict[str, float]) -> pd.DataFrame:
    """
    Sum the Transaction amounts by Client, converting to USD using fx_rates.
    Transactions in currencies not present in fx_rates are excluded from the total.
    """
    df = df.copy()
    df["Transaction"] = pd.to_numeric(df["Transaction"], errors="coerce")
    df["Exchange_Rate"] = df["Currency"].map(fx_rates)
    df["Transaction_USD"] = df["Transaction"] * df["Exchange_Rate"]
    return df.groupby("Client")["Transaction_USD"].sum().reset_index(name="Total_USD")


############################################################################################
#                           Internal analysis helpers                                      #
############################################################################################

def _enrich(df_raw: pd.DataFrame, fx_rates: dict[str, float]) -> pd.DataFrame:
    """Add computed columns used across all analysis functions."""
    df = df_raw.copy()
    df["Transaction"]     = pd.to_numeric(df["Transaction"], errors="coerce")
    df["Exchange_Rate"]   = df["Currency"].map(fx_rates)
    df["Transaction_USD"] = df["Transaction"] * df["Exchange_Rate"]
    df["Market_Section"]  = df["Client"].str[0].str.upper()
    df["Month"]           = df["log_date"].dt.to_period("M")
    return df


def _cc_summary(df: pd.DataFrame) -> pd.DataFrame:
    """Client-Currency pairs: transaction count, local total, USD total."""
    return (
        df.groupby(["Client", "Currency"])
        .agg(
            Transaction_Count=("Transaction", "count"),
            Total_Local=("Transaction", "sum"),
            Total_USD=("Transaction_USD", "sum"),
        )
        .reset_index()
        .sort_values("Total_USD", ascending=False)
        .reset_index(drop=True)
    )


def _client_summary(df: pd.DataFrame) -> pd.DataFrame:
    """Per-client totals with rank, section, and transaction count."""
    totals = (
        df.groupby("Client")["Transaction_USD"]
        .sum()
        .reset_index(name="Total_USD")
        .sort_values("Total_USD", ascending=False)
        .reset_index(drop=True)
    )
    totals["Market_Section"]    = totals["Client"].str[0].str.upper()
    totals["Transaction_Count"] = df.groupby("Client").size().reindex(totals["Client"]).values
    totals.insert(0, "Rank", range(1, len(totals) + 1))
    return totals


def _monthly_summary(df: pd.DataFrame) -> pd.DataFrame:
    """Revenue and transaction count by calendar month."""
    monthly = (
        df.groupby("Month")
        .agg(Transaction_Count=("Transaction", "count"), Total_USD=("Transaction_USD", "sum"))
        .reset_index()
    )
    monthly["Month_Label"] = monthly["Month"].dt.strftime("%b %Y")
    return monthly


def _section_summary(df: pd.DataFrame) -> pd.DataFrame:
    """Revenue by market section (first character of client code)."""
    section = (
        df.groupby("Market_Section")
        .agg(
            Client_Count=("Client", "nunique"),
            Transaction_Count=("Transaction", "count"),
            Total_USD=("Transaction_USD", "sum"),
        )
        .reset_index()
        .sort_values("Total_USD", ascending=False)
        .reset_index(drop=True)
    )
    section["Avg_per_Client_USD"] = section["Total_USD"] / section["Client_Count"]
    return section


def _geo_summary(df: pd.DataFrame) -> pd.DataFrame:
    """Revenue and transaction count by country."""
    return (
        df.groupby("Country")
        .agg(Transaction_Count=("Transaction", "count"), Total_USD=("Transaction_USD", "sum"))
        .reset_index()
        .sort_values("Total_USD", ascending=False)
        .reset_index(drop=True)
    )


############################################################################################
#                           Excel generation                                               #
############################################################################################

_XL_NAVY  = "1C3F6E"
_XL_BLUE  = "2E75B6"
_XL_LIGHT = "BDD7EE"
_XL_ALT   = "EBF3FB"
_XL_WHITE = "FFFFFF"
_XL_DARK  = "1F2D3D"


def _fill(h):
    return PatternFill("solid", fgColor=h)


def _font(c=_XL_DARK, bold=False, sz=10, nm="Calibri"):
    return Font(name=nm, size=sz, bold=bold, color=c)


def _border():
    s = Side(style="thin", color="D0D0D0")
    return Border(left=s, right=s, top=s, bottom=s)


def _write_table(ws, df, sr=1, sc=1, hdr=_XL_NAVY):
    """Fully styled per-cell write — use only for small tables (< ~500 rows)."""
    b = _border()
    ws.row_dimensions[sr].height = 22
    for c, col in enumerate(df.columns, start=sc):
        cell = ws.cell(row=sr, column=c, value=col)
        cell.fill = _fill(hdr); cell.font = _font(_XL_WHITE, True, 10)
        cell.border = b
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(df.itertuples(index=False), start=sr + 1):
        rf = _fill(_XL_ALT) if (r - sr) % 2 == 0 else _fill(_XL_WHITE)
        for c, val in enumerate(row, start=sc):
            cell = ws.cell(row=r, column=c, value=val)
            cell.fill = rf; cell.border = b; cell.font = _font()
            if isinstance(val, float):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif isinstance(val, int):
                cell.number_format = "#,##0"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")


def _write_fast(ws, df, col_widths=None):
    """Header-only styling + ws.append() for large tables. col_widths: list of ints."""
    b = _border()
    ws.row_dimensions[1].height = 22
    for c, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=c, value=col)
        cell.fill = _fill(_XL_NAVY); cell.font = _font(_XL_WHITE, True, 10)
        cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = b
    for row in df.itertuples(index=False):
        ws.append(list(row))
    if col_widths:
        for i, w in enumerate(col_widths, start=1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w


def _aw(ws, lo=8, hi=45):
    for col in ws.columns:
        w = max((len(str(c.value or "")) for c in col), default=lo)
        ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, lo), hi)


def _build_excel(total_rev, total_txn, n_clients, n_countries,
                 cc, client_summary, monthly, section, geo,
                 output_path, year_label="FY 2024"):
    wb = openpyxl.Workbook()

    # ── Sheet 1: Executive Summary ────────────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Executive Summary"
    ws_sum.sheet_view.showGridLines = False

    for rh, val, sz, bold in [
        (1, "",  8,  False),
        (2, f"MARSH |  FP&A Regional Analysis — {year_label}", 20, True),
        (3, "Comprehensive Transaction & Revenue Report", 12, False),
        (4, "", 8, False),
    ]:
        ws_sum.merge_cells(f"A{rh}:J{rh}")
        cell = ws_sum.cell(row=rh, column=1, value=val)
        cell.fill = _fill(_XL_NAVY)
        cell.font = Font(name="Calibri", size=sz, bold=bold, color=_XL_WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws_sum.row_dimensions[rh].height = 8 if not val else (48 if sz > 15 else 22)

    ws_sum.row_dimensions[6].height = 50
    ws_sum.row_dimensions[7].height = 20
    ws_sum.row_dimensions[8].height = 8

    for value, label, c1, c2 in [
        (f"${total_rev / 1e6:.1f} M", "Total Revenue (USD)",  "B", "C"),
        (f"{total_txn:,}",             "Total Transactions",   "D", "E"),
        (f"{n_clients:,}",             "Unique Clients",       "F", "G"),
        (f"{n_countries:,}",           "Countries",            "H", "I"),
    ]:
        ws_sum.merge_cells(f"{c1}6:{c2}6")
        ws_sum.merge_cells(f"{c1}7:{c2}7")
        v = ws_sum[f"{c1}6"]
        v.value = value; v.fill = _fill(_XL_BLUE)
        v.font = Font(name="Calibri", size=22, bold=True, color=_XL_WHITE)
        v.alignment = Alignment(horizontal="center", vertical="center")
        l = ws_sum[f"{c1}7"]
        l.value = label; l.fill = _fill(_XL_LIGHT)
        l.font = Font(name="Calibri", size=10, bold=True, color=_XL_NAVY)
        l.alignment = Alignment(horizontal="center", vertical="center")

    ws_sum.merge_cells("A10:J10")
    h = ws_sum["A10"]
    h.value = "  REPORT CONTENTS"; h.fill = _fill(_XL_NAVY)
    h.font = _font(_XL_WHITE, True, 11)
    h.alignment = Alignment(horizontal="left", vertical="center")
    ws_sum.row_dimensions[10].height = 22

    for i, (sn, desc) in enumerate([
        ("Txns by Client & Currency",  "Count and total amount per client-currency pair"),
        ("Client Revenue (USD)",        "All clients ranked by USD revenue with embedded chart"),
        ("Monthly Revenue",             "Revenue and transaction volume by month with line chart"),
        ("Market Section Analysis",     "Revenue by market section (first character of client code)"),
        ("Geographic Analysis",         "Revenue ranked by country with bar chart"),
    ], start=11):
        ws_sum.row_dimensions[i].height = 18
        rf = _fill(_XL_ALT) if i % 2 == 0 else _fill(_XL_WHITE)
        for c in range(1, 11):
            ws_sum.cell(row=i, column=c).fill = rf
        ws_sum.cell(row=i, column=1, value=f"  {i - 10}.").font = _font(_XL_DARK, sz=10)
        n = ws_sum.cell(row=i, column=2, value=sn)
        n.font = Font(name="Calibri", size=10, bold=True, color=_XL_BLUE, underline="single")
        ws_sum.merge_cells(f"C{i}:J{i}")
        ws_sum.cell(row=i, column=3, value=desc).font = _font()

    ws_sum.column_dimensions["A"].width = 5
    ws_sum.column_dimensions["B"].width = 36
    for c in ["C", "D", "E", "F", "G", "H", "I", "J"]:
        ws_sum.column_dimensions[c].width = 14

    # ── Sheet 2: Txns by Client & Currency  (large — fast path) ──────────────
    ws_cc = wb.create_sheet("Txns by Client & Currency")
    ws_cc.sheet_view.showGridLines = False
    ws_cc.freeze_panes = "A2"
    cc_e = cc.copy()
    cc_e.columns = ["Client", "Currency", "Transaction Count",
                    "Total Amount (Local)", "Total Amount (USD)"]
    _write_fast(ws_cc, cc_e, col_widths=[14, 12, 18, 22, 22])

    # ── Sheet 3: Client Revenue (USD)  (large — fast path) ───────────────────
    ws_cli = wb.create_sheet("Client Revenue (USD)")
    ws_cli.sheet_view.showGridLines = False
    ws_cli.freeze_panes = "A2"
    cli_e = client_summary[["Rank", "Client", "Total_USD",
                             "Transaction_Count", "Market_Section"]].copy()
    cli_e.columns = ["Rank", "Client", "Total Revenue (USD)",
                     "Transaction Count", "Market Section"]
    _write_fast(ws_cli, cli_e, col_widths=[8, 14, 22, 18, 16])

    ch = BarChart()
    ch.type = "bar"; ch.title = f"Top 20 Clients by Revenue (USD)"; ch.style = 10
    ch.y_axis.title = "Client"; ch.x_axis.title = "Revenue (USD)"
    ch.width = 22; ch.height = 16
    tn = min(21, len(cli_e) + 1)
    ch.add_data(Reference(ws_cli, min_col=3, min_row=1, max_row=tn), titles_from_data=True)
    ch.set_categories(Reference(ws_cli, min_col=2, min_row=2, max_row=tn))
    ws_cli.add_chart(ch, "G2")

    # ── Sheet 4: Monthly Revenue ──────────────────────────────────────────────
    ws_mon = wb.create_sheet("Monthly Revenue")
    ws_mon.sheet_view.showGridLines = False
    mon_e = monthly[["Month_Label", "Transaction_Count", "Total_USD"]].copy()
    mon_e.columns = ["Month", "Transaction Count", "Total Revenue (USD)"]
    _write_table(ws_mon, mon_e)
    for r in range(2, len(mon_e) + 2):
        ws_mon.cell(r, 2).number_format = "#,##0"
        ws_mon.cell(r, 3).number_format = '"$"#,##0.00'
    _aw(ws_mon)

    ch2 = LineChart()
    ch2.title = f"Monthly Revenue — {year_label}"; ch2.style = 10
    ch2.y_axis.title = "Revenue (USD)"; ch2.x_axis.title = "Month"
    ch2.width = 22; ch2.height = 14
    nm = len(mon_e)
    ch2.add_data(Reference(ws_mon, min_col=3, min_row=1, max_row=nm + 1), titles_from_data=True)
    ch2.set_categories(Reference(ws_mon, min_col=1, min_row=2, max_row=nm + 1))
    ws_mon.add_chart(ch2, "E2")

    # ── Sheet 5: Market Section Analysis ─────────────────────────────────────
    ws_sec = wb.create_sheet("Market Section Analysis")
    ws_sec.sheet_view.showGridLines = False
    sec_e = section.copy()
    sec_e.columns = ["Market Section", "Unique Clients", "Transaction Count",
                     "Total Revenue (USD)", "Avg Revenue / Client (USD)"]
    _write_table(ws_sec, sec_e)
    for r in range(2, len(sec_e) + 2):
        ws_sec.cell(r, 2).number_format = "#,##0"
        ws_sec.cell(r, 3).number_format = "#,##0"
        ws_sec.cell(r, 4).number_format = '"$"#,##0.00'
        ws_sec.cell(r, 5).number_format = '"$"#,##0.00'
    _aw(ws_sec)

    ch3 = BarChart()
    ch3.title = "Total Revenue by Market Section"; ch3.style = 10
    ch3.y_axis.title = "Revenue (USD)"; ch3.x_axis.title = "Market Section"
    ch3.width = 20; ch3.height = 14
    ns = len(sec_e)
    ch3.add_data(Reference(ws_sec, min_col=4, min_row=1, max_row=ns + 1), titles_from_data=True)
    ch3.set_categories(Reference(ws_sec, min_col=1, min_row=2, max_row=ns + 1))
    ws_sec.add_chart(ch3, "G2")

    # ── Sheet 6: Geographic Analysis ─────────────────────────────────────────
    ws_geo = wb.create_sheet("Geographic Analysis")
    ws_geo.sheet_view.showGridLines = False
    ws_geo.freeze_panes = "A2"
    geo_e = geo.copy()
    geo_e.insert(0, "Rank", range(1, len(geo_e) + 1))
    geo_e.columns = ["Rank", "Country", "Transaction Count", "Total Revenue (USD)"]
    _write_table(ws_geo, geo_e)
    for r in range(2, len(geo_e) + 2):
        ws_geo.cell(r, 3).number_format = "#,##0"
        ws_geo.cell(r, 4).number_format = '"$"#,##0.00'
    _aw(ws_geo)

    ch4 = BarChart()
    ch4.type = "bar"; ch4.title = "Top 20 Countries by Revenue (USD)"; ch4.style = 10
    ch4.y_axis.title = "Country"; ch4.x_axis.title = "Revenue (USD)"
    ch4.width = 20; ch4.height = 15
    tg = min(21, len(geo_e) + 1)
    ch4.add_data(Reference(ws_geo, min_col=4, min_row=1, max_row=tg), titles_from_data=True)
    ch4.set_categories(Reference(ws_geo, min_col=2, min_row=2, max_row=tg))
    ws_geo.add_chart(ch4, "F2")

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)


############################################################################################
#                           Main entry point                                               #
############################################################################################

def generate_report(
    data_root: str,
    fx_path: str | None = None,
    output_path: str | None = None,
    progress_cb=None,
) -> str:
    """
    Full pipeline: load → enrich → analyse → export Excel.

    Parameters
    ----------
    data_root   : root folder containing the DataTrans_* files (may be nested).
    fx_path     : path to the FX Rates xlsx; if None, the first match for
                  *FX*Rates*.xlsx found recursively under data_root is used.
    output_path : destination .xlsx path; defaults to
                  <data_root>/Deliverable_1_FPA_Analysis.xlsx.
    progress_cb : optional callable(message: str) for GUI progress updates.

    Returns the absolute path to the saved workbook.
    """
    def _log(msg):
        if progress_cb:
            progress_cb(msg)
        else:
            print(msg)

    # Locate FX file if not provided
    if fx_path is None:
        hits = glob.glob(os.path.join(data_root, "**", "*FX*Rates*.xlsx"), recursive=True)
        if not hits:
            hits = glob.glob(os.path.join(data_root, "**", "*FX*.xlsx"), recursive=True)
        if not hits:
            raise FileNotFoundError(
                f"No FX rates file found under '{data_root}'. "
                "Pass fx_path explicitly."
            )
        fx_path = hits[0]

    if output_path is None:
        output_path = os.path.join(data_root, "Deliverable_1_FPA_Analysis.xlsx")

    _log(f"Loading transactions from: {data_root}")
    df_raw = load_all_transactions(data_root)
    if df_raw.empty:
        raise ValueError(f"No transaction files found under '{data_root}'.")
    _log(f"  {len(df_raw):,} rows loaded")

    _log(f"Loading FX rates from: {fx_path}")
    fx_rates = load_fx_rates(fx_path)

    _log("Enriching data...")
    df = _enrich(df_raw, fx_rates)

    # Derive year label from the actual data date range
    years = df["log_date"].dt.year.dropna().unique()
    if len(years) == 1:
        year_label = f"FY {int(years[0])}"
    elif len(years) > 1:
        year_label = f"FY {int(years.min())}–{int(years.max())}"
    else:
        year_label = "FY Report"

    total_rev   = df["Transaction_USD"].sum()
    total_txn   = len(df)
    n_clients   = df["Client"].nunique()
    n_countries = df["Country"].nunique()

    _log("Running analyses...")
    cc             = _cc_summary(df)
    client_summary = _client_summary(df)
    monthly        = _monthly_summary(df)
    section        = _section_summary(df)
    geo            = _geo_summary(df)

    _log(f"Writing Excel report to: {output_path}")
    _build_excel(
        total_rev, total_txn, n_clients, n_countries,
        cc, client_summary, monthly, section, geo,
        output_path, year_label,
    )

    out = os.path.abspath(output_path)
    _log(f"Done → {out}")
    return out
